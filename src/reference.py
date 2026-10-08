"""Read reference metadata separately from competition telemetry."""
from pathlib import Path
import json
import re

import pandas as pd
from openpyxl import load_workbook

from src.config import ANALYSIS_VERSION, SOURCE_CONTRACT_VERSION, EVENT_REFERENCE_VERSION
from src.utils import sha256, write_json, markdown_table

EVENT_FILE = 'Event Type_0929補充.xlsx'
FAQ_FILE = 'HINO competition FAQ.pdf'
RAW_SHEETS = ['14-Rawdata(已遷移)', '14V2-Rawdata(已遷移)', '14V3-Rawdata(已遷移)']
EXTENDED_FIELDS = {
    'driverUid': ('駕駛 ID', '取代 Prototype 模擬駕駛代碼；正式提供情況待確認'),
    'tachographDriver': ('大餅駕駛 ID', '核對駕駛輪替；正式提供情況待確認'),
    'engineFuelRate': ('引擎燃油率', '更直接估算怠速與區段燃油；單位及可用性待確認'),
    'instantaneousFuel': ('瞬時燃油', '輔助區段燃油分析；單位及可用性待確認'),
    'pedalPosition': ('油門踏板位置', '理解操作需求；可用性待確認'),
    'gearBoxPosition': ('變速箱檔位', '理解轉速與檔位型態；可用性待確認'),
    'ptoSwitch': ('PTO 開關', '判斷停留時是否可能有必要作業；可用性待確認'),
    'streetName': ('道路名稱', '增加路段情境；可用性待確認'),
    'streetLevel': ('道路層級', '改善路段分組；可用性待確認'),
    'brakeSwitch': ('煞車開關', '理解操作情境；可用性待確認'),
    'clutchSwitch': ('離合器開關', '理解操作情境；可用性待確認'),
}


def build_reference(root: Path, competition_columns, competition_sha=None):
    root = Path(root)
    path = root / 'data/reference' / EVENT_FILE
    if not path.exists():
        raise FileNotFoundError(f'Event reference missing: {path}')
    book = load_workbook(path, read_only=True, data_only=True)
    event_rows = []
    last = None
    for row in book['event代碼'].iter_rows(min_row=2, values_only=True):
        code, name, onoff, definition = row[:4]
        if isinstance(code, (int, float)):
            last = int(code)
            event_rows.append({'event_type': last, 'event_name_zh': str(name or ''),
                               'trigger_definition': str(definition or ''), 'parameter_configurable': last in {2, 6, 7, 8, 11},
                               'configuration_dependent': last in {2, 6, 7, 8, 11},
                               'comparison_role': 'SUPPORTING_ONLY' if last in {2, 6, 7, 8, 11} else 'REFERENCE_ONLY',
                               'source_note': f'{EVENT_FILE} / event代碼；原表開關={onoff}',
                               'source_conflict_flag': last == 8})
        elif last is not None and name:
            event_rows[-1]['trigger_definition'] += f'；{name}：{definition or ""}'
    events = pd.DataFrame(event_rows)
    events.to_parquet(root / 'data/reference/event_reference.parquet', index=False)
    records = []
    for sheet in RAW_SHEETS:
        rows = book[sheet].iter_rows(values_only=True)
        for row in rows:
            if not row or len(row) < 2 or not isinstance(row[1], str):
                continue
            name = row[1].strip()
            if not re.fullmatch(r'[A-Za-z][A-Za-z0-9._\[\]-]*', name) or name in {'info', 'type'}:
                continue
            records.append({'field': name, 'display_name': str(row[0] or ''), 'source_sheet': sheet,
                            'data_type': str(row[2] or '') if len(row) > 2 else '',
                            'reference_note': str(row[7] or '') if len(row) > 7 else ''})
    fields = pd.DataFrame(records).drop_duplicates(['field', 'source_sheet'])
    fields['competition_available'] = fields.field.isin(competition_columns)
    fields['production_availability'] = 'TO_CONFIRM'
    fields.to_parquet(root / 'data/reference/extended_itraq_fields.parquet', index=False)
    matrix = []
    for field in ['enabledCode', 'journeyCode', 'Type', 'gps.longitude', 'gps.latitude', 'can.totalMileage',
                  'can.engine.totalFuelUsed', 'can.engine.rpm', 'can.engine.engineLoad', 'carStatus', 'event[0].type']:
        matrix.append(dict(field=field, display_name=field, data_layer='A_COMPETITION_CORE',
                           competition_available=field in competition_columns, reference_schema_documented=False,
                           production_availability='AVAILABLE_IN_COMPETITION', current_prototype_use='RAW / DERIVED',
                           future_product_use='可比較行程證據', limitations='本次 extract 實際欄位與單位以競賽活頁簿為準',
                           source_file='output data_Hotai_20260511.xlsx', source_sheet='output / 欄位對照表'))
    for field, (label, use) in EXTENDED_FIELDS.items():
        matches = fields.loc[fields.field.eq(field)]
        matrix.append(dict(field=field, display_name=label, data_layer='B_EXTENDED_ITRAQ',
                           competition_available=field in competition_columns,
                           reference_schema_documented=not matches.empty,
                           production_availability='PRODUCTION_AVAILABILITY_TO_CONFIRM',
                           current_prototype_use='NONE', future_product_use=use,
                           limitations='參考 schema 已遷移；依車型、設備及管線確認。不得視為本次競賽實測資料。',
                           source_file=EVENT_FILE,
                           source_sheet=' / '.join(matches.source_sheet.tolist())))
    for field, use in [('task_id', '確認任務'), ('actual_payload', '控制實際載重'), ('delivery_window', '核對到站窗口'),
                       ('stop_purpose', '區分停留用途'), ('loading_status', '核對裝卸作業'), ('maintenance_history', '核對保養')]:
        matrix.append(dict(field=field, display_name=field, data_layer='C_ENTERPRISE_CONTEXT', competition_available=False,
                           reference_schema_documented=False, production_availability='ENTERPRISE_INTEGRATION_REQUIRED',
                           current_prototype_use='MOCK placeholder' if field in {'stop_purpose', 'delivery_window'} else 'NONE',
                           future_product_use=use, limitations='需企業營運系統串接', source_file='', source_sheet=''))
    for field in ['traffic', 'weather', 'road_gradient']:
        matrix.append(dict(field=field, display_name=field, data_layer='D_EXTERNAL_CONTEXT', competition_available=False,
                           reference_schema_documented=False, production_availability='EXTERNAL_OPTIONAL',
                           current_prototype_use='MOCK placeholder' if field == 'traffic' else 'NONE',
                           future_product_use='擴充環境比較條件', limitations='須取得外部可靠來源', source_file='', source_sheet=''))
    capability = pd.DataFrame(matrix)
    capability.to_csv(root / 'data/reference/data_capability_matrix.csv', index=False)
    registry = {'analysis_version': ANALYSIS_VERSION, 'source_contract_version': SOURCE_CONTRACT_VERSION,
                'event_reference_version': EVENT_REFERENCE_VERSION,
                'sources': [
                    {'priority': 'A', 'file': 'output data_Hotai_20260511.xlsx', 'role': 'competition observations and units',
                     'sha256': competition_sha},
                    {'priority': 'B', 'file': FAQ_FILE, 'role': 'competition definitions and limitations',
                     'sha256': sha256(root / 'data/reference' / FAQ_FILE) if (root / 'data/reference' / FAQ_FILE).exists() else None},
                    {'priority': 'C', 'file': EVENT_FILE, 'sheet': 'event代碼', 'role': 'event metadata only', 'sha256': sha256(path)},
                    {'priority': 'D', 'file': EVENT_FILE, 'sheets': RAW_SHEETS, 'role': 'extended production schema reference only'},
                    {'priority': 'E', 'file': EVENT_FILE, 'sheet': '14.1-胎壓資料格式(已作廢)', 'role': 'deprecated; excluded'}],
                'conflicts': [{'code': 'SOURCE_CONFLICT', 'subject': 'Event 8 speeding threshold',
                               'resolution': 'Retain source-reported Event 8; do not reconstruct thresholds.'},
                              {'code': 'REFERENCE_CONFLICT', 'subject': 'event duration unit',
                               'competition_unit': 'seconds', 'reference_unit': 'minutes in migrated 14V3-Rawdata rows',
                               'resolution': 'Competition extract and its 欄位對照表 take precedence; idle_minutes never uses event duration.'}]}
    write_json(root / 'data/reference/source_registry.json', registry)
    selected = events.loc[events.event_type.isin([2, 6, 7, 8, 11])].copy()
    selected['source_conflict_flag'] = selected.source_conflict_flag.map({True: 'SOURCE_CONFLICT', False: ''})
    event_doc = '# iTRAQ 事件參考\n\n來源：`Event Type_0929補充.xlsx` 的 `event代碼` sheet；僅為 reference metadata，不是新增 telemetry。\n\n'
    event_doc += markdown_table(selected.to_dict('records'), ['event_type', 'event_name_zh', 'trigger_definition',
        'parameter_configurable', 'comparison_role', 'source_conflict_flag', 'source_note'])
    event_doc += ('\n\nEvent 2 的觸發／解除門檻不等於完整怠速時長。`idle_minutes` 只依 `carStatus=idling` 與有效 timestamp interval 計算。\n\n'
                  'Event 6／7／8／11 為設定依賴的來源事件，僅作支持證據；事件計數不是跨車標準化績效。'
                  'Event 8 競賽 FAQ 第 2–3 頁與 0929 的預設門檻不同，標記 `SOURCE_CONFLICT`；保留來源系統通報，不自行重建。'
                  '`engine_load_above_90_ratio` 是 EcoPilot 原型特徵，不等於來源 Event 11（參考描述 >83% 持續 5 秒，參數可調）。\n')
    (root / 'docs/event_reference.md').write_text(event_doc, encoding='utf-8')
    matrix_doc = '# 四層資料能力矩陣\n\nA 為此次 Competition Core；B 為已遷移 iTRAQ Rawdata 的參考 schema，正式可用性需核對；C 需企業營運系統；D 為選配外部資料。Mock 是 C／D 的 Prototype placeholder。\n\n'
    matrix_doc += markdown_table(capability.to_dict('records'), list(capability.columns))
    matrix_doc += '\n\nB 層不代表競賽 extract 已觀測到這些值，也不代表所有車型或設備都提供；欄位不進入 telemetry 或主要模型。已作廢胎壓表僅作 deprecated reference。\n'
    (root / 'docs/data_capability_matrix.md').write_text(matrix_doc, encoding='utf-8')
    return events, fields, capability, registry
