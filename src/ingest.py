"""Read-only, bounded-chunk Excel ingestion; profile the entire dataset."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

from src.utils import ROOT, markdown_table, sha256, write_json
from src.config import TYPE_DEFINITION

WORKBOOK_NAME = "output data_Hotai_20260511.xlsx"


def find_workbook(root=ROOT):
    root = Path(root)
    direct = root / "data/raw" / WORKBOOK_NAME
    if direct.exists():
        return direct.resolve()
    candidates = sorted(p for p in root.parent.rglob(WORKBOOK_NAME) if ".venv" not in p.parts)
    if not candidates:
        raise FileNotFoundError(f"請將 {WORKBOOK_NAME} 放入 {root / 'data/raw'}")
    if len(candidates) > 1:
        raise ValueError("找到多個 workbook，請使用 --input 指定：" + ", ".join(map(str, candidates)))
    return candidates[0].resolve()


def _text(value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def inspect_workbook(path, output_dir, chunk_size=25000):
    path, output_dir = Path(path), Path(output_dir)
    interim, docs = output_dir / "data/interim", output_dir / "docs"
    interim.mkdir(parents=True, exist_ok=True)
    docs.mkdir(parents=True, exist_ok=True)
    source_hash = sha256(path)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = [{"sheet": s.title, "rows_including_header": s.max_row, "columns": s.max_column} for s in workbook]
    definitions, mapping_rows = {}, []
    if "欄位對照表" in workbook.sheetnames:
        for row in workbook["欄位對照表"].iter_rows(min_row=2, values_only=True):
            if len(row) >= 7 and row[2] is not None:
                d = {"description": _text(row[1]), "field": str(row[2]).strip(), "declared_type": _text(row[3]),
                     "range": _text(row[4]), "unit": _text(row[5]) or "UNKNOWN", "example": _text(row[6])}
                mapping_rows.append(d)
                definitions.setdefault(d["field"], d)
    sheet = workbook["output"] if "output" in workbook.sheetnames else workbook.worksheets[0]
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(v).strip() if v is not None else f"unnamed_{i}" for i, v in enumerate(next(iterator))]
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate workbook column names; inspect source headers before ingestion")
    schema = pa.schema([(h, pa.string()) for h in headers] + [("source_row", pa.int64())])
    target = interim / "raw_snapshot.parquet"
    temporary = target.with_suffix(".parquet.tmp")
    row_count = 0
    with pq.ParquetWriter(temporary, schema, compression="zstd") as writer:
        batch = []
        for source_row, row in enumerate(iterator, start=2):
            record = {h: _text(v) for h, v in zip(headers, row)}
            if not any(record.values()):
                continue
            record["source_row"] = source_row
            batch.append(record)
            row_count += 1
            if len(batch) >= chunk_size:
                writer.write_table(pa.Table.from_pylist(batch, schema=schema))
                print(f"Read {row_count:,} rows", flush=True)
                batch.clear()
        if batch:
            writer.write_table(pa.Table.from_pylist(batch, schema=schema))
    workbook.close()
    temporary.replace(target)
    frame = pd.read_parquet(target)
    profile = profile_frame(frame, headers, definitions)
    profile.update({"source_file": path.name, "source_sha256": source_hash, "sheets": sheets,
                    "telemetry_sheet": sheet.title, "row_count": row_count, "column_count": len(headers),
                    "field_definitions": definitions})
    if sha256(path) != source_hash:
        raise RuntimeError("Source workbook checksum changed during ingestion")
    write_json(docs / "data_profile.json", profile)
    write_json(docs / "workbook_field_mapping.json", mapping_rows)
    write_json(interim / "source_manifest.json", {"source_path": str(path.resolve()), "sha256": source_hash,
               "rows": row_count, "columns": headers, "ingestion_version": 1})
    write_profile_markdown(profile, docs / "data_profile.md")
    return profile


def profile_frame(frame, headers, definitions):
    p = {"columns": {}, "distributions": {}, "numeric_summaries": {}, "quality": {}}
    for col in headers:
        s = frame[col]
        nonnull = s.dropna()
        numeric = pd.to_numeric(nonnull, errors="coerce")
        dtype = "EMPTY" if nonnull.empty else "number" if numeric.notna().all() else "string"
        if col == "time" or col.endswith("Time") and "event[" in col:
            parsed = pd.to_datetime(nonnull, errors="coerce", format="mixed")
            dtype = "datetime" if parsed.notna().all() else "mixed datetime/string"
        if col in ["enabledCode", "journeyCode", "carNum"]:
            dtype = "identifier (stored as string)"
        definition = definitions.get(col, definitions.get(col.split(".info.")[-1], {}))
        p["columns"][col] = {"inferred_type": dtype, "missing_pct": round(float(s.isna().mean() * 100), 4),
                            "unique_count": int(nonnull.nunique()), "unit": definition.get("unit", "UNKNOWN"),
                            "description": definition.get("description", "UNKNOWN")}
        if col in ["Type", "carStatus", "can.canStatus"] or col.endswith(".type"):
            p["distributions"][col] = s.fillna("MISSING").value_counts().to_dict()
        if col in ["can.totalMileage", "can.engine.totalFuelUsed", "line.fuelLevel", "can.engine.engineTotalTime",
                    "can.engine.rpm", "can.engine.engineLoad", "gps.speed", "can.canSpeed"]:
            p["numeric_summaries"][col] = numeric.describe(percentiles=[.01, .5, .95, .99]).to_dict()
    p["vehicle_count"] = int(frame["enabledCode"].nunique()) if "enabledCode" in frame else 0
    p["journey_count"] = int(frame["journeyCode"].nunique()) if "journeyCode" in frame else 0
    key = ["enabledCode", "journeyCode"]
    p["trip_count"] = len(frame[key].dropna().drop_duplicates()) if set(key) <= set(frame) else 0
    if "time" in frame:
        t = pd.to_datetime(frame["time"], errors="coerce", format="mixed")
        p["time_range"] = [str(t.min()), str(t.max())]
        p["quality"]["invalid_timestamps"] = int(t.isna().sum())
        if set(key) <= set(frame):
            ordered = frame.assign(timestamp=t).sort_values(key + ["timestamp", "source_row"])
            gaps = ordered.groupby(key).timestamp.diff().dt.total_seconds().dropna()
            p["sampling_interval_seconds"] = gaps.describe(percentiles=[.01, .25, .5, .75, .9, .95, .99]).to_dict()
            p["sampling_interval_top_counts"] = gaps.value_counts().head(15).to_dict()
            p["quality"]["duplicate_trip_timestamps"] = int(ordered.duplicated(key + ["timestamp"]).sum())
            p["quality"]["gaps_over_300s"] = int((gaps > 300).sum())
            p["quality"]["journey_codes_shared_across_vehicles"] = int((frame.groupby("journeyCode").enabledCode.nunique() > 1).sum())
            p["quality"]["exact_duplicate_records"] = int(frame[headers].duplicated().sum())
            p["quality"]["missing_trip_key_rows"] = int(frame[key].isna().any(axis=1).sum())
            for col in ["can.totalMileage", "can.engine.totalFuelUsed", "can.engine.engineTotalTime"]:
                if col in ordered:
                    delta = pd.to_numeric(ordered[col], errors="coerce").groupby([ordered[k] for k in key]).diff()
                    p["quality"][col + "_negative_deltas"] = int((delta < 0).sum())
                    p["quality"][col + "_positive_delta_summary"] = delta[delta > 0].describe().to_dict()
    if {"gps.speed", "can.canSpeed"} <= set(frame):
        gps, can = (pd.to_numeric(frame[c], errors="coerce") for c in ["gps.speed", "can.canSpeed"])
        p["quality"]["gps_zero_when_can_moving_rows"] = int(((gps == 0) & (can > 5)).sum())
    return p


def write_profile_markdown(p, path):
    sections = ["# HINO iTRAQ — 全量資料檢查", f"原始檔：`{p['source_file']}`；SHA-256：`{p['source_sha256']}`。",
                "使用 openpyxl 唯讀模式分批掃描全資料，並非抽樣。所有空白字串視為缺值；原始 xlsx 保持不變。",
                "## 活頁簿工作表", markdown_table([{"工作表": s["sheet"], "含標題列數": s["rows_including_header"], "欄位數": s["columns"]} for s in p["sheets"]], ["工作表", "含標題列數", "欄位數"]),
                f"車聯網資料：**{p['row_count']:,} 列 × {p['column_count']} 欄**。",
                f"車輛數：{p['vehicle_count']}；不重複 journeyCode 數：{p['journey_count']}；複合行程鍵數：{p['trip_count']}。",
                "時間範圍：" + " → ".join(p.get("time_range", ["UNKNOWN"])),
                "## 欄位、推定型態、缺值與單位", "單位只取自活頁簿欄位對照表；未註明者一律標為未知。" + TYPE_DEFINITION,
                markdown_table([{"原始欄位代碼": k, "推定型態": {"number": "數值", "datetime": "日期時間", "string": "文字", "mixed datetime/string": "日期時間／文字混合", "identifier (stored as string)": "識別碼（以文字儲存）", "EMPTY": "全空"}.get(v.get("inferred_type"), v.get("inferred_type")), "缺值比例": v.get("missing_pct"), "不重複值數": v.get("unique_count"), "單位": "未註明" if v.get("unit") == "UNKNOWN" else v.get("unit"), "說明": "未註明" if v.get("description") == "UNKNOWN" else v.get("description")} for k, v in p["columns"].items()],
                               ["原始欄位代碼", "推定型態", "缺值比例", "不重複值數", "單位", "說明"]),
                "以下統計區塊保留原始欄位代碼與 JSON 鍵名，供技術查核。"]
    for title, field in [("Type、狀態與事件分布", "distributions"), ("燃油、里程與數值欄位統計", "numeric_summaries"),
                         ("時間戳記採樣間隔（秒，依同一複合行程鍵排序）", "sampling_interval_seconds"),
                         ("常見間隔（秒 → 次數）", "sampling_interval_top_counts"), ("資料品質檢查", "quality")]:
        sections += ["## " + title, "```json\n" + json.dumps(p.get(field, {}), ensure_ascii=False, indent=2, default=str) + "\n```"]
    sections += ["## 累積欄位、可疑欄位與限制",
                 "- `can.totalMileage`（公里）、`can.engine.totalFuelUsed`（公升）、`can.engine.engineTotalTime`（小時）為累積欄位；負差分須標記計數器重設。`line.fuelLevel` 是油箱百分比，不能當作累積耗油。",
                 "- `journeyCode` 可能跨車重複，只能與 `enabledCode` 組成複合行程鍵。`carNum` 不作車輛識別碼。",
                 "- 原始狀態代碼 0／1／2 分別代表停車、行駛、怠速。GPS 與 CAN 速度矛盾時須標記，不自行捏造真實狀態。",
                 "- GPS 座標在合法範圍內仍可能失準。重複座標、GPS 零速而 CAN 顯示移動，以及軌跡跳點，都會限制路線判讀。",
                 "- 原始欄位對照表的引擎負載上限為 250%，不能把超過 100% 的值一律視為錯誤；事件中的負載欄位也可能有不同定義。",
                 "- 原始時間欄位沒有時區資訊；以 Asia/Taipei 解讀營運時段和班次，是明示的原型假設。",
                 "- 事件可能重複通報。事件長表保留每筆原始內容；計數時依事件識別條件去重。",
                 "- 真實載重、駕駛、等待目的、天氣與交通狀況無法從此活頁簿確認，只能在原型中明確標記為模擬資料。"]
    Path(path).write_text("\n\n".join(sections) + "\n", encoding="utf-8")
